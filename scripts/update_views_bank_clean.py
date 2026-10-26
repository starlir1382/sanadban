with open('workflow/views.py', 'r', encoding='utf-8') as f:
    views_code = f.read()

# 1. Update models import
if 'ProjectBankMapping' not in views_code:
    views_code = views_code.replace(
        'from .models import HousingDossier',
        'from .models import HousingDossier, ProjectBankMapping'
    )

views_norm = views_code.replace('\r\n', '\n')
old_start = "@login_required(login_url='login')\ndef bank_letter_sender_view(request):"
pos_start = views_norm.find(old_start)
assert pos_start != -1, "old_start not found"

new_section = """@login_required(login_url='login')
def project_bank_mapping_api(request):
    \"\"\"
    API to fetch and update Project-to-Bank mappings for all 96 projects.
    \"\"\"
    profile = getattr(request.user, 'profile', None)
    is_chief_or_admin = request.user.is_superuser or request.user.is_staff or (profile and profile.is_housing_chief)
    can_action_bank = profile and (profile.can_action_stage(6) or profile.can_action_stage(9))
    if not (is_chief_or_admin or can_action_bank):
        return JsonResponse({"status": "error", "message": "دسترسی غیرمجاز."}, status=403)

    if request.method == "POST":
        try:
            body = json.loads(request.body.decode("utf-8")) if request.body else {}
        except Exception:
            body = request.POST

        mappings = body.get("mappings")
        if mappings and isinstance(mappings, list):
            updated = 0
            for item in mappings:
                p = (item.get("project_name") or "").strip()
                b = (item.get("bank_code") or "maskan").strip()
                if p and b in ["maskan", "melli", "mellat", "saderat", "tejarat"]:
                    ProjectBankMapping.objects.update_or_create(
                        project_name=p,
                        defaults={"bank_code": b}
                    )
                    updated += 1
            return JsonResponse({"status": "success", "message": f"تخصیص {updated} پروژه به بانک‌های عامل با موفقیت ذخیره شد."})

        project_name = (body.get("project_name") or "").strip()
        bank_code = (body.get("bank_code") or "").strip()
        if not project_name or bank_code not in ["maskan", "melli", "mellat", "saderat", "tejarat"]:
            return JsonResponse({"status": "error", "message": "اطلاعات نامعتبر است."}, status=400)

        mapping_obj, _ = ProjectBankMapping.objects.update_or_create(
            project_name=project_name,
            defaults={"bank_code": bank_code}
        )
        return JsonResponse({
            "status": "success",
            "message": f"پروژه «{project_name}» به {mapping_obj.get_bank_code_display()} تخصیص داده شد.",
            "project_name": project_name,
            "bank_code": bank_code,
            "bank_name": mapping_obj.get_bank_code_display()
        })

    # GET: return all 96 projects mappings
    from django.db.models import Count
    all_projects = sorted(list(set(p.strip() for p in HousingDossier.objects.values_list("project", flat=True) if p and p.strip())))
    existing_maps = {m.project_name: m.bank_code for m in ProjectBankMapping.objects.all()}
    proj_counts = dict(HousingDossier.objects.values("project").annotate(c=Count("id")).values_list("project", "c"))

    result = []
    for p in all_projects:
        b_code = existing_maps.get(p, "maskan")
        result.append({
            "project_name": p,
            "bank_code": b_code,
            "bank_name": ProjectBankMapping.BANK_NAMES.get(b_code, "بانک مسکن"),
            "dossier_count": proj_counts.get(p, 0)
        })

    return JsonResponse({"status": "success", "projects": result, "total": len(result)})


@login_required(login_url='login')
def bank_letter_sender_view(request):
    \"\"\"
    Admin component to prepare, preview and send official Bank Introduction Letter batches.
    Displays the list of eligible dossiers (who reach full payment through loans + cash).
    Supports multi-project sheet export for 5 banks (Melli, Mellat, Maskan, Saderat, Tejarat).
    \"\"\"
    profile = getattr(request.user, 'profile', None)
    is_chief_or_admin = request.user.is_superuser or request.user.is_staff or (profile and profile.is_housing_chief)
    can_action_bank = profile and (profile.can_action_stage(6) or profile.can_action_stage(9))
    
    if not (is_chief_or_admin or can_action_bank):
        messages.error(request, "دسترسی به بخش ارسال نامه بانک نیازمند مجوز کارشناس بانک یا مدیر سیستم است.")
        return redirect("dossier_list")

    # All available dossiers
    all_dossiers = HousingDossier.objects.all().order_by("archive_number")
    
    # 96 Distinct Projects
    projects = sorted(list(set(p.strip() for p in HousingDossier.objects.values_list("project", flat=True) if p and p.strip())))

    # Project-Bank Mappings
    existing_maps = {m.project_name: m.bank_code for m in ProjectBankMapping.objects.all()}
    # Auto-initialize any unmapped project to maskan
    for p in projects:
        if p not in existing_maps:
            ProjectBankMapping.objects.get_or_create(project_name=p, defaults={"bank_code": "maskan"})
            existing_maps[p] = "maskan"

    # Default today Jalali date
    try:
        import jdatetime
        today_jalali = jdatetime.date.today().strftime("%Y/%m/%d")
    except Exception:
        today_jalali = timezone.now().strftime("%Y/%m/%d")

    # Sync any sent dossiers into BankLetter if not already present
    sent_dossiers = HousingDossier.objects.filter(is_sent_to_bank=True).exclude(bank_letter_number="")
    for d in sent_dossiers:
        num = (d.bank_letter_number or d.bank_sent_letter_number or "").strip()
        if num:
            bl_obj, _ = BankLetter.objects.get_or_create(
                letter_number=num,
                defaults={
                    "letter_date": d.bank_letter_date or d.bank_sent_date or today_jalali,
                    "project_name": d.project or "پروژه مسکونی",
                    "sender_name": "کارشناس تسهیلات بانکی",
                    "sender_role": "واحد تسهیلات و بانک مسکن",
                }
            )
            if not bl_obj.dossiers.filter(id=d.id).exists():
                bl_obj.dossiers.add(d)

    # Fetch Letter History
    letters_qs = BankLetter.objects.all().prefetch_related("dossiers").order_by("-created_at", "-id")
    letter_history_list = []
    for l in letters_qs:
        d_objs = list(l.dossiers.all().order_by("archive_number"))
        try:
            import jdatetime
            dt_jalali = jdatetime.datetime.fromgregorian(datetime=l.created_at).strftime("%Y/%m/%d - %H:%M")
        except Exception:
            dt_jalali = l.created_at.strftime("%Y/%m/%d %H:%M") if l.created_at else ""

        letter_history_list.append({
            "id": l.id,
            "letter_number": l.letter_number,
            "letter_date": l.letter_date,
            "project_name": l.project_name or (d_objs[0].project if d_objs else "پروژه مسکونی"),
            "sender_name": l.sender_name or "کارشناس تسهیلات",
            "sender_role": l.sender_role or "واحد تسهیلات و بانک مسکن",
            "notes": l.notes or "",
            "created_at_jalali": dt_jalali,
            "dossier_count": len(d_objs),
            "dossier_archives": [d.archive_number for d in d_objs],
            "dossiers": [
                {
                    "archive_number": d.archive_number,
                    "full_name": d.full_name or "-",
                    "national_id": d.national_id or "-",
                    "block_code": d.block_code or d.block_name or "-",
                    "floor_number": d.floor_number or "-",
                    "unit_code": d.unit_code or d.unit_number or "-",
                    "total_area": str(d.total_area or "-"),
                    "applicant_phone": d.applicant_phone or "-",
                }
                for d in d_objs
            ]
        })

    # Bank Statistics & Project Mappings JSON
    from django.db.models import Count
    proj_dossier_counts = dict(HousingDossier.objects.values("project").annotate(c=Count("id")).values_list("project", "c"))
    
    bank_stats = {
        code: {
            "code": code,
            "name": name,
            "project_count": 0,
            "dossier_count": 0
        }
        for code, name in ProjectBankMapping.BANK_CHOICES
    }

    project_mappings_list = []
    for p in projects:
        b_code = existing_maps.get(p, "maskan")
        d_count = proj_dossier_counts.get(p, 0)
        if b_code in bank_stats:
            bank_stats[b_code]["project_count"] += 1
            bank_stats[b_code]["dossier_count"] += d_count
        project_mappings_list.append({
            "project_name": p,
            "bank_code": b_code,
            "bank_name": ProjectBankMapping.BANK_NAMES.get(b_code, "بانک مسکن"),
            "dossier_count": d_count
        })

    # Serialize dossiers for Alpine.js store
    dossiers_json_list = []
    fully_funded_count = 0
    for d in all_dossiers:
        fn, ln = split_persian_name(d.full_name)
        base_cost = d.base_cost_without_loan or 0
        total_loans = (d.fund_loan_amount or 0) + (d.decile_loan_amount or 0)
        is_funded = (total_loans >= base_cost and base_cost > 0) or d.is_equity_funded or d.is_sent_to_bank or (d.current_stage >= 6)
        if is_funded:
            fully_funded_count += 1

        b_code = existing_maps.get(d.project, "maskan")

        dossiers_json_list.append({
            "id": d.id,
            "archive_number": d.archive_number,
            "national_id": d.national_id or "",
            "first_name": fn,
            "last_name": ln,
            "full_name": d.full_name or "",
            "selection_rank": d.selection_rank or "-",
            "applicant_phone": d.applicant_phone or "",
            "block_code": d.block_code or d.block_name or "-",
            "floor_number": d.floor_number or "-",
            "unit_code": d.unit_code or d.unit_number or "-",
            "covered_parking_area": str(d.covered_parking_area or "-"),
            "parking_number": str(d.parking_number or "-"),
            "storage_area": str(d.storage_area or "-"),
            "storage_number": str(d.storage_number or "-"),
            "bedroom_count": str(d.bedroom_count or "-"),
            "common_area": str(d.common_area or "-"),
            "net_unit_area": str(d.net_unit_area or "-"),
            "unit_direction": d.unit_direction or "-",
            "total_area": str(d.total_area or "-"),
            "project": d.project or "",
            "site": d.site or "",
            "bank_code": b_code,
            "bank_name": ProjectBankMapping.BANK_NAMES.get(b_code, "بانک مسکن"),
            "current_stage": d.current_stage,
            "current_stage_title": d.current_stage_title,
            "is_sent_to_bank": d.is_sent_to_bank,
            "is_fully_funded": is_funded,
            "base_cost_without_loan": base_cost,
            "total_loans": total_loans,
            "applicant_equity": d.applicant_equity,
            "is_equity_funded": d.is_equity_funded,
            "bank_letter_number": d.bank_letter_number or d.bank_sent_letter_number or "",
            "bank_letter_date": d.bank_letter_date or d.bank_sent_date or "",
        })

    initial_letter = request.GET.get("letter", "").strip()

    context = {
        "dossiers_json": json.dumps(dossiers_json_list, ensure_ascii=False),
        "letters_history_json": json.dumps(letter_history_list, ensure_ascii=False),
        "project_mappings_json": json.dumps(project_mappings_list, ensure_ascii=False),
        "bank_stats": bank_stats,
        "banks_list": ProjectBankMapping.BANK_CHOICES,
        "letters_count": len(letter_history_list),
        "projects": projects,
        "projects_count": len(projects),
        "total_dossiers_count": len(dossiers_json_list),
        "fully_funded_count": fully_funded_count,
        "today_jalali": today_jalali,
        "default_project": projects[0] if projects else "پروژه مسکونی",
        "initial_letter": initial_letter,
    }
    return render(request, "admin/bank_letter_sender.html", context)


@login_required(login_url='login')
def bank_letter_submit_api(request):
    \"\"\"
    Processes the batch Bank Letter submission:
    Updates is_sent_to_bank, letter number and date, creates BankLetter archive entity,
    and advances Stage 6 dossiers to Stage 7.
    \"\"\"
    profile = getattr(request.user, 'profile', None)
    is_chief_or_admin = request.user.is_superuser or request.user.is_staff or (profile and profile.is_housing_chief)
    can_action_bank = profile and (profile.can_action_stage(6) or profile.can_action_stage(9))
    
    if not (is_chief_or_admin or can_action_bank):
        return JsonResponse({"status": "error", "message": "دسترسی غیرمجاز."}, status=403)

    try:
        body = json.loads(request.body.decode("utf-8")) if request.body else {}
    except Exception:
        body = request.POST

    archive_numbers = body.get("archive_numbers", [])
    letter_number = (body.get("letter_number") or "").strip()
    letter_date = (body.get("letter_date") or "").strip()
    notes = (body.get("notes") or "").strip()

    if not archive_numbers:
        return JsonResponse({"status": "error", "message": "هیچ پرونده‌ای انتخاب نشده است."}, status=400)
    if not letter_number:
        return JsonResponse({"status": "error", "message": "شماره نامه بانک الزامی است."}, status=400)
    if not letter_date:
        return JsonResponse({"status": "error", "message": "تاریخ نامه بانک الزامی است."}, status=400)

    clean_archives = []
    for item in archive_numbers:
        try:
            clean_archives.append(int(str(item).strip()))
        except ValueError:
            continue

    dossiers = HousingDossier.objects.filter(archive_number__in=clean_archives)
    updated_count = 0
    updated_archives = []

    actor_name = profile.full_name if profile and profile.full_name else request.user.username
    actor_role = profile.role_title if profile and profile.role_title else "کارشناس تسهیلات بانکی"

    for d in dossiers:
        d.is_sent_to_bank = True
        d.bank_letter_number = letter_number
        d.bank_letter_date = letter_date
        d.bank_sent_letter_number = letter_number
        d.bank_sent_date = letter_date
        d.is_added_to_bank_shipment_list = True
        d.bank_shipment_list_added_at = timezone.now()

        from_stage_str = f"مرحله {d.current_stage}: {d.current_stage_title}"
        if d.current_stage == 6:
            d.current_stage = 7
            to_stage_str = f"مرحله ۷: {d.current_stage_title}"
        else:
            to_stage_str = f"مرحله {d.current_stage}: {d.current_stage_title}"

        d.save()
        updated_count += 1
        updated_archives.append(d.archive_number)

        WorkflowSignatureLog.objects.create(
            dossier=d,
            from_stage=from_stage_str,
            to_stage=to_stage_str,
            actor=actor_name,
            role=actor_role,
            action="BANK_LETTER",
            action_label="ارسال نامه معرفی به بانک",
            notes=f"نامه معرفی به بانک صادره به شماره {letter_number} مورخ {letter_date}. {notes}".strip(),
            signature_hash=f"BANK-LETTER-{d.archive_number}-{letter_number}"
        )

    project_title = dossiers.first().project if dossiers.exists() and dossiers.first().project else "پروژه مسکونی"
    bank_letter, created = BankLetter.objects.get_or_create(
        letter_number=letter_number,
        defaults={
            "letter_date": letter_date,
            "project_name": project_title,
            "sender_user": request.user,
            "sender_name": actor_name,
            "sender_role": actor_role,
            "notes": notes,
        }
    )
    bank_letter.letter_date = letter_date
    if notes:
        bank_letter.notes = notes
    bank_letter.sender_name = actor_name
    bank_letter.sender_role = actor_role
    bank_letter.dossiers.add(*dossiers)
    bank_letter.save()

    return JsonResponse({
        "status": "success",
        "message": f"نامه بانک با شماره {letter_number} برای {updated_count} پرونده با موفقیت ثبت و وضعیت پرونده‌ها به‌روزرسانی شد.",
        "updated_count": updated_count,
        "updated_archives": updated_archives,
        "letter_number": letter_number,
        "letter_date": letter_date,
        "letter_id": bank_letter.id
    })


def _build_bank_workbook(bank_code, bank_name, project_names, dossiers_qs):
    \"\"\"
    Builds an openpyxl Workbook where each project assigned to this bank gets its own sheet.
    Exact 18-column standard format from official project unit sheet.
    \"\"\"
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    # Remove default sheet
    default_sheet = wb.active
    wb.remove(default_sheet)

    thin_border = Border(
        left=Side(style='thin', color='475569'),
        right=Side(style='thin', color='475569'),
        top=Side(style='thin', color='475569'),
        bottom=Side(style='thin', color='475569')
    )
    headers = [
        "ردیف", "نام", "نام خانوادگی", "رتبه", "کدملی", "شماره تماس",
        "کد بلوک", "طبقه", "کدواحد", "متراژ پارکینگ مسقف", "شماره پارکینگ",
        "متراژ انباری", "شماره انباری", "تعداد خواب", "مشاعات", "متراژ مفید",
        "جهت", "جمع متراژ"
    ]
    col_widths = [8, 16, 18, 8, 14, 15, 12, 8, 10, 16, 14, 14, 14, 12, 12, 12, 12, 12]

    sheet_names_used = set()
    for proj in project_names:
        clean_name = re.sub(r'[:\\\\/?*\\[\\]]', '-', proj).strip()
        short_name = clean_name[:30] if len(clean_name) > 30 else clean_name
        if not short_name:
            short_name = "پروژه"
        base_name = short_name
        counter = 1
        while short_name in sheet_names_used:
            short_name = f"{base_name[:26]}_{counter}"
            counter += 1
        sheet_names_used.add(short_name)

        ws = wb.create_sheet(title=short_name)
        ws.views.sheetView[0].rightToLeft = True

        # Title Banner (Row 1)
        ws.merge_cells("A1:R1")
        title_cell = ws["A1"]
        title_cell.value = f"محاسبه قیمت تمام شده انتخاب واحد پروژه {proj} ({bank_name})"
        title_cell.font = Font(name="Tahoma", size=12, bold=True, color="000000")
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        title_cell.fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
        ws.row_dimensions[1].height = 36

        # Header Row (Row 2)
        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=2, column=col_idx, value=h)
            cell.font = Font(name="Tahoma", size=10, bold=True, color="000000")
            cell.fill = PatternFill(start_color="CBD5E1", end_color="CBD5E1", fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border
        ws.row_dimensions[2].height = 28

        # Project dossiers
        proj_dossiers = dossiers_qs.filter(project=proj).order_by("archive_number")
        row_idx = 3
        for idx, d in enumerate(proj_dossiers, 1):
            fn, ln = split_persian_name(d.full_name)
            row_vals = [
                idx,
                fn or d.full_name,
                ln,
                d.selection_rank or "",
                d.national_id or "",
                d.applicant_phone or "",
                d.block_code or d.block_name or "",
                d.floor_number or "",
                d.unit_code or d.unit_number or "",
                float(d.covered_parking_area) if d.covered_parking_area else "",
                d.parking_number or "",
                float(d.storage_area) if d.storage_area else "",
                d.storage_number or "",
                d.bedroom_count or "",
                float(d.common_area) if d.common_area else "",
                float(d.net_unit_area) if d.net_unit_area else "",
                d.unit_direction or "",
                float(d.total_area) if d.total_area else "",
            ]
            for col_idx, val in enumerate(row_vals, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.font = Font(name="Tahoma", size=9)
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = thin_border
            ws.row_dimensions[row_idx].height = 22
            row_idx += 1

        for col_idx, w in enumerate(col_widths, 1):
            col_letter = get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = w

    if not wb.sheetnames:
        ws = wb.create_sheet(title=bank_name[:30])
        ws.views.sheetView[0].rightToLeft = True
        ws["A1"] = f"پروژه‌ای برای {bank_name} ثبت نگردیده است."

    return wb


@login_required(login_url='login')
def bank_letter_export_excel(request):
    \"\"\"
    Exports Bank Introduction Letters:
    1) ?bank=melli / mellat / maskan / saderat / tejarat: Exports an Excel with a sheet for each project of that bank.
    2) ?bank=zip: Exports all 5 banks Excels packed in a single ZIP file.
    3) default: Single-sheet export for filtered archives or letter.
    \"\"\"
    import io
    import urllib.parse
    import zipfile
    from django.http import HttpResponse

    bank_req = request.GET.get("bank", "").strip().lower()
    archive_str = request.GET.get("archives", "")
    letter_number = request.GET.get("letter", "").strip()
    letter_date = request.GET.get("date", "").strip()
    only_funded = request.GET.get("only_funded", "0") == "1"

    # Base queryset
    if archive_str:
        arch_list = [int(a.strip()) for a in archive_str.split(",") if a.strip().isdigit()]
        dossiers_qs = HousingDossier.objects.filter(archive_number__in=arch_list)
    elif letter_number:
        dossiers_qs = HousingDossier.objects.filter(bank_letter_number=letter_number)
    else:
        dossiers_qs = HousingDossier.objects.all()

    # Project to bank map
    existing_maps = {m.project_name: m.bank_code for m in ProjectBankMapping.objects.all()}
    all_projects = sorted(list(set(p.strip() for p in HousingDossier.objects.values_list("project", flat=True) if p and p.strip())))

    # 1. Multi-bank ZIP Export: All 5 Banks Excels in one ZIP
    if bank_req == "zip":
        zip_buffer = io.BytesIO()
        with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zip_file:
            for b_code, b_name in ProjectBankMapping.BANK_CHOICES:
                b_projs = [p for p in all_projects if existing_maps.get(p, "maskan") == b_code]
                wb = _build_bank_workbook(b_code, b_name, b_projs, dossiers_qs)
                excel_buffer = io.BytesIO()
                wb.save(excel_buffer)
                excel_buffer.seek(0)
                file_name = f"Nameh_Bank_{b_name.replace(' ', '_')}.xlsx"
                zip_file.writestr(file_name, excel_buffer.getvalue())

        zip_buffer.seek(0)
        resp = HttpResponse(zip_buffer.getvalue(), content_type="application/zip")
        resp["Content-Disposition"] = 'attachment; filename="5_Banks_Project_Excel_Letters.zip"'
        return resp

    # 2. Specific Bank Export (e.g. maskan, melli, mellat, saderat, tejarat)
    valid_banks = dict(ProjectBankMapping.BANK_CHOICES)
    if bank_req in valid_banks:
        b_name = valid_banks[bank_req]
        b_projs = [p for p in all_projects if existing_maps.get(p, "maskan") == bank_req]
        wb = _build_bank_workbook(bank_req, b_name, b_projs, dossiers_qs)
        excel_buffer = io.BytesIO()
        wb.save(excel_buffer)
        excel_buffer.seek(0)
        resp = HttpResponse(excel_buffer.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        enc_fname = urllib.parse.quote(f"Nameh_Bank_{b_name.replace(' ', '_')}.xlsx")
        resp["Content-Disposition"] = f'attachment; filename="{enc_fname}"; filename*=UTF-8\'\'{enc_fname}'
        return resp

    # 3. Default single-sheet export
    project_title = request.GET.get("project", "پروژه مسکونی")
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "نامه بانک"
    ws.views.sheetView[0].rightToLeft = True

    sheet_title = f"محاسبه قیمت تمام شده انتخاب واحد پروژه {project_title}"
    if letter_number:
        sheet_title += f" (نامه رسمی بانک: {letter_number} مورخ {letter_date})"

    ws.merge_cells("A1:R1")
    title_cell = ws["A1"]
    title_cell.value = sheet_title
    title_cell.font = Font(name="Tahoma", size=13, bold=True, color="000000")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    title_cell.fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
    ws.row_dimensions[1].height = 36

    headers = [
        "ردیف", "نام", "نام خانوادگی", "رتبه", "کدملی", "شماره تماس",
        "کد بلوک", "طبقه", "کدواحد", "متراژ پارکینگ مسقف", "شماره پارکینگ",
        "متراژ انباری", "شماره انباری", "تعداد خواب", "مشاعات", "متراژ مفید",
        "جهت", "جمع متراژ"
    ]
    thin_border = Border(
        left=Side(style='thin', color='475569'),
        right=Side(style='thin', color='475569'),
        top=Side(style='thin', color='475569'),
        bottom=Side(style='thin', color='475569')
    )
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=2, column=col_idx, value=h)
        cell.font = Font(name="Tahoma", size=10, bold=True, color="000000")
        cell.fill = PatternFill(start_color="CBD5E1", end_color="CBD5E1", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    ws.row_dimensions[2].height = 28

    row_idx = 3
    for idx, d in enumerate(dossiers_qs.order_by("archive_number"), 1):
        fn, ln = split_persian_name(d.full_name)
        row_vals = [
            idx,
            fn or d.full_name,
            ln,
            d.selection_rank or "",
            d.national_id or "",
            d.applicant_phone or "",
            d.block_code or d.block_name or "",
            d.floor_number or "",
            d.unit_code or d.unit_number or "",
            float(d.covered_parking_area) if d.covered_parking_area else "",
            d.parking_number or "",
            float(d.storage_area) if d.storage_area else "",
            d.storage_number or "",
            d.bedroom_count or "",
            float(d.common_area) if d.common_area else "",
            float(d.net_unit_area) if d.net_unit_area else "",
            d.unit_direction or "",
            float(d.total_area) if d.total_area else "",
        ]
        for col_idx, val in enumerate(row_vals, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = Font(name="Tahoma", size=9)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border
        ws.row_dimensions[row_idx].height = 22
        row_idx += 1

    col_widths = [8, 16, 18, 8, 14, 15, 12, 8, 10, 16, 14, 14, 14, 12, 12, 12, 12, 12]
    for col_idx, w in enumerate(col_widths, 1):
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = w

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    resp = HttpResponse(output.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = 'attachment; filename="Nameh_Bank.xlsx"'
    return resp
"""

views_norm = views_norm[:pos_start] + new_section

if '\r\n' in views_code:
    views_out = views_norm.replace('\n', '\r\n')
else:
    views_out = views_norm

with open('workflow/views.py', 'w', encoding='utf-8') as f:
    f.write(views_out)

print("views.py updated successfully with 5-bank exports and project mapping API")
